# -*- coding: utf-8 -*-
"""Оформление документов отчётности ФСИ по образцу сданной отчётности проекта «Опти».

Профили:
  report — отчёт о выполнении работ: TNR 12, интервал 1,5, без титульного листа, нумерация с 2;
  bp     — бизнес-план: TNR 14 (таблицы 12), титульный лист без номера;
  dev    — отчёт о развитии проекта: TNR 12, одна таблица;
  app    — приложения к отчёту: как report, нумерация страниц с 1.
"""
import os
import re
from docx import Document
from docx.shared import Pt, Mm, Cm, RGBColor
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER,
                            WD_COLOR_INDEX)
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from PIL import Image, ImageDraw, ImageFont

A = WD_ALIGN_PARAGRAPH
NBSP = " "
DASH = "–"     # короткое тире: подписи «Таблица 1 – …», диапазоны
MDASH = "—"    # длинное тире в тексте и в перечне сокращений («АУСН — …»)
FONT = "Times New Roman"
TEXT_W_MM = 165
TEXT_W_TW = 9355
FONT_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


class Profile:
    def __init__(self, body, table, header_fill, list_left, list_hanging, h2_center, h2_caps,
                 tcap_before, tcap_after, fig_before, title_page, pg_start, toc_l2_indent_cm):
        self.body = body
        self.table = table
        self.header_fill = header_fill
        self.list_left = list_left
        self.list_hanging = list_hanging
        self.h2_center = h2_center
        self.h2_caps = h2_caps
        self.tcap_before = tcap_before
        self.tcap_after = tcap_after
        self.fig_before = fig_before
        self.title_page = title_page
        self.pg_start = pg_start
        self.toc_l2_indent_cm = toc_l2_indent_cm


PROFILES = {
    "report": Profile(body=12, table=12, header_fill="D9E1F2", list_left=1134, list_hanging=425,
                      h2_center=True, h2_caps=True, tcap_before=9, tcap_after=3, fig_before=9,
                      title_page=False, pg_start=2, toc_l2_indent_cm=0.0),
    "bp": Profile(body=14, table=12, header_fill="EDEDED", list_left=1069, list_hanging=360,
                  h2_center=False, h2_caps=False, tcap_before=12, tcap_after=3, fig_before=6,
                  title_page=True, pg_start=None, toc_l2_indent_cm=0.75),
    "dev": Profile(body=12, table=12, header_fill="D9D9D9", list_left=1134, list_hanging=425,
                   h2_center=False, h2_caps=False, tcap_before=6, tcap_after=3, fig_before=6,
                   title_page=False, pg_start=None, toc_l2_indent_cm=0.0),
    # приложения к отчету — отдельные файлы в оформлении отчета, нумерация страниц с 1
    "app": Profile(body=12, table=12, header_fill="D9E1F2", list_left=1134, list_hanging=425,
                   h2_center=True, h2_caps=True, tcap_before=9, tcap_after=3, fig_before=9,
                   title_page=False, pg_start=None, toc_l2_indent_cm=0.0),
}


def nb(s):
    return str(s).replace(" ", NBSP)


def fmt_num(x, dec=0):
    """Число по-русски: разряды через неразрывный пробел, десятичная запятая, минус «−»."""
    neg = x < 0
    x = abs(x)
    if dec == 0:
        s = f"{int(round(x)):,}".replace(",", NBSP)
    else:
        s = f"{x:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", NBSP)
    return ("−" if neg and s.strip("0, ") else "") + s


# ------------------------------------------------------------------ порядок элементов OOXML
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
ORDER_SEQ = {
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
              "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
              "tblLook", "tblCaption", "tblDescription", "tblPrChange"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
             "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel", "cellMerge",
             "tcPrChange"],
    "sectPr": ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
               "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
               "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange"],
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
            "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
            "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
            "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"],
}


def _local(tag):
    return tag.split("}")[1] if isinstance(tag, str) and "}" in tag else tag


def _normalize_element(el, seq, dedupe):
    if dedupe:
        seen = set()
        for c in [c for c in el if isinstance(c.tag, str)]:
            name = _local(c.tag)
            if name in seen:
                el.remove(c)
            else:
                seen.add(name)
    kids = [c for c in el if isinstance(c.tag, str)]
    rank = {n: i for i, n in enumerate(seq)}
    ordered = sorted(kids, key=lambda c: rank.get(_local(c.tag), len(seq)))
    if [id(c) for c in ordered] != [id(c) for c in kids]:
        for c in kids:
            el.remove(c)
        for c in ordered:
            el.append(c)


def normalize_order(doc):
    for part in doc.part.package.iter_parts():
        el = getattr(part, "_element", None)
        if el is None or not str(part.partname).startswith("/word/"):
            continue
        for tag, seq in ORDER_SEQ.items():
            for node in list(el.iter(_W + tag)):
                if tag == "rPr" and node.getparent() is not None and _local(node.getparent().tag) == "rPrChange":
                    continue
                _normalize_element(node, seq, dedupe=tag in ("tblPr", "tcPr"))


_UNITS = r"(?:°C|%|руб\.|руб\b|тыс\.|млн|млрд|г\b|кг\b|мм\b|см\b|мл\b|мкм\b|шт\.|чел\.|дней\b|дня\b|суток\b|мин\b|с\b|м\b|px\b)"


def typo(text):
    """Неразрывные пробелы: между числом и единицей, в разрядах, после «№», в «г. Москва», в инициалах."""
    t = re.sub(r"(\d) (?=" + _UNITS + r")", "\\1" + NBSP, text)
    t = re.sub(r"(?<=\d) (?=\d{3}\b)", NBSP, t)
    t = t.replace("№ ", "№" + NBSP)
    t = re.sub(r"\bг\. (?=[А-ЯЁ])", "г." + NBSP, t)
    t = re.sub(r"\b([А-ЯЁ])\. ([А-ЯЁ])\. (?=[А-ЯЁ][а-яё])", "\\1." + NBSP + "\\2." + NBSP, t)
    return t


# ------------------------------------------------------------------ разметка текста
_MARK_RE = re.compile(r"(\[УТОЧНИТЬ[^\]]*\]|\*\*.+?\*\*)")


class Doc:
    def __init__(self, profile, title="", resolver=None):
        self.P = PROFILES[profile]
        self.kind = profile
        self.doc = Document()
        body = self.doc.element.body
        for p in body.findall(qn("w:p")):
            body.remove(p)
        self.R = resolver or (lambda t: t)
        self._defaults()
        self._styles()
        self._page()
        self._numbering()
        cp = self.doc.core_properties
        cp.title = title
        cp.author = "ООО «Вкус Технологии»"
        cp.last_modified_by = "ООО «Вкус Технологии»"
        cp.subject = ""
        cp.keywords = ""
        cp.comments = ""
        cp.category = ""

    # ---------------------------------------------------------------- базовые настройки
    def _defaults(self):
        st = self.doc.styles.element
        dd = st.find(qn("w:docDefaults"))
        if dd is not None:
            st.remove(dd)
        sz = self.P.body * 2
        dd = parse_xml(
            f'<w:docDefaults {nsdecls("w")}><w:rPrDefault><w:rPr>'
            f'<w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:eastAsia="{FONT}" w:cs="{FONT}"/>'
            f'<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/><w:lang w:val="ru-RU" w:eastAsia="ru-RU" w:bidi="ar-SA"/>'
            f'</w:rPr></w:rPrDefault><w:pPrDefault><w:pPr><w:spacing w:after="0" w:line="360" w:lineRule="auto"/>'
            f'</w:pPr></w:pPrDefault></w:docDefaults>')
        st.insert(0, dd)

    def _style_font(self, s, size, bold):
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.italic = False
        s.font.color.rgb = RGBColor(0, 0, 0)
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None:
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), FONT)
        for tag in ("w:szCs",):
            el = rpr.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                rpr.append(el)
            el.set(qn("w:val"), str(size * 2))

    def _style_par(self, s, align, first_cm, before, after, keep_next=True):
        pf = s.paragraph_format
        pf.alignment = align
        pf.first_line_indent = Cm(first_cm)
        pf.left_indent = Cm(0)
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.5
        pf.keep_with_next = keep_next
        pf.keep_together = False

    def _styles(self):
        P = self.P
        n = self.doc.styles["Normal"]
        self._style_font(n, P.body, None)
        n.font.bold = None
        pf = n.paragraph_format
        pf.alignment = A.JUSTIFY
        pf.first_line_indent = Cm(1.25)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.5
        pf.widow_control = True
        h1 = self.doc.styles["Heading 1"]
        self._style_font(h1, P.body, True)
        self._style_par(h1, A.CENTER, 0, 0 if self.kind == "bp" else 18, 12)
        h2 = self.doc.styles["Heading 2"]
        self._style_font(h2, P.body, True)
        if P.h2_center:
            self._style_par(h2, A.CENTER, 0, 12, 6)
        else:
            self._style_par(h2, A.JUSTIFY, 1.25, 12, 6)
        h3 = self.doc.styles["Heading 3"]
        self._style_font(h3, P.body, True)
        self._style_par(h3, A.JUSTIFY, 1.25, 6, 3)
        for s in (h1, h2, h3):
            ppr = s.element.get_or_add_pPr()
            kl = ppr.find(qn("w:keepLines"))
            if kl is not None:
                ppr.remove(kl)
        lp = self.doc.styles["List Paragraph"]
        ppr = lp.element.get_or_add_pPr()
        for tag in ("w:ind", "w:contextualSpacing"):
            el = ppr.find(qn(tag))
            if el is not None:
                ppr.remove(el)

    def _page(self):
        P = self.P
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin, sec.right_margin = Mm(30), Mm(15)
        sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)
        sec.header_distance = Mm(12.5)
        sec.footer_distance = Mm(10)
        sec.different_first_page_header_footer = P.title_page
        fp = sec.footer.paragraphs[0]
        fp.alignment = A.CENTER
        fp.paragraph_format.first_line_indent = Cm(0)
        fp.paragraph_format.line_spacing = 1.0
        fld = parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" '
            f'w:cs="{FONT}"/><w:sz w:val="{P.body * 2}"/><w:szCs w:val="{P.body * 2}"/></w:rPr><w:t>2</w:t></w:r></w:fldSimple>')
        fp._p.append(fld)
        sp = sec._sectPr
        pg = sp.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            sp.append(pg)
        if P.pg_start:
            pg.set(qn("w:start"), str(P.pg_start))

    def _numbering(self):
        self.numbering = self.doc.part.numbering_part.element
        P = self.P
        sz = P.body * 2

        def add_abs(aid, fmt, text):
            xml = (f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}"><w:multiLevelType w:val="singleLevel"/>'
                   f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/>'
                   f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="{P.list_left}" w:hanging="{P.list_hanging}"/></w:pPr>'
                   f'<w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:cs="{FONT}"/><w:sz w:val="{sz}"/>'
                   f'<w:szCs w:val="{sz}"/></w:rPr></w:lvl></w:abstractNum>')
            el = parse_xml(xml)
            first_num = self.numbering.find(qn("w:num"))
            if first_num is not None:
                first_num.addprevious(el)
            else:
                self.numbering.append(el)

        add_abs(90, "bullet", "–")
        add_abs(91, "decimal", "%1)")
        self.num_bullet = self._add_num(90)

    def _add_num(self, aid, restart=False):
        ids = [int(n.get(qn("w:numId"))) for n in self.numbering.findall(qn("w:num"))]
        nid = max(ids + [100]) + 1
        xml = f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="{aid}"/>'
        if restart:
            xml += '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>'
        xml += "</w:num>"
        self.numbering.append(parse_xml(xml))
        return nid

    # ---------------------------------------------------------------- текст
    def _runs(self, par, text, size=None, bold=None, italic=None):
        text = typo(self.R(text))
        size = size or self.P.body
        for part in _MARK_RE.split(text):
            if not part:
                continue
            b = bold
            hl = False
            if part.startswith("[УТОЧНИТЬ"):
                hl = True
            elif part.startswith("**") and part.endswith("**") and len(part) > 4:
                part = part[2:-2]
                b = True
            r = par.add_run(part)
            r.font.name = FONT
            r.font.size = Pt(size)
            rpr = r._element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), FONT)
            if b is not None:
                r.font.bold = b
            if italic is not None:
                r.font.italic = italic
            if hl:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return par

    def _pf(self, par, align=A.JUSTIFY, first=True, before=0, after=0, line=1.5, keep_next=False,
            page_break_before=False, left_cm=None):
        pf = par.paragraph_format
        pf.alignment = align
        pf.first_line_indent = Cm(1.25) if first is True else Cm(first or 0)
        if left_cm is not None:
            pf.left_indent = Cm(left_cm)
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line
        pf.keep_with_next = keep_next
        if page_break_before:
            pf.page_break_before = True
        return par

    def p(self, text, align=A.JUSTIFY, first=True, bold=None, italic=None, size=None, before=0, after=0,
          line=1.5, keep_next=False, page_break_before=False, left_cm=None):
        par = self.doc.add_paragraph()
        self._pf(par, align, first, before, after, line, keep_next, page_break_before, left_cm)
        self._runs(par, text, size, bold, italic)
        return par

    def empty(self, line=1.0, size=None):
        par = self.doc.add_paragraph()
        self._pf(par, A.LEFT, 0, 0, 0, line)
        r = par.add_run("")
        r.font.size = Pt(size or self.P.body)
        return par

    # ---------------------------------------------------------------- заголовки
    def struct_title(self, text, page_break=True):
        """Заголовок структурного элемента вне оглавления (РЕФЕРАТ, СОДЕРЖАНИЕ): по центру, жирный."""
        par = self.doc.add_paragraph()
        self._pf(par, A.CENTER, 0, 0 if self.kind == "bp" else 18, 12, 1.5, True, page_break)
        self._runs(par, text.upper(), bold=True)
        return par

    def h1(self, text, page_break=True):
        par = self.doc.add_paragraph(style="Heading 1")
        if page_break:
            par.paragraph_format.page_break_before = True
        self._runs(par, text.upper())
        return par

    def h2(self, text, page_break=False, first_in_section=False):
        par = self.doc.add_paragraph(style="Heading 2")
        if page_break:
            par.paragraph_format.page_break_before = True
        if first_in_section and not self.P.h2_center:
            par.paragraph_format.space_before = Pt(6)
        self._runs(par, text.upper() if self.P.h2_caps else text)
        return par

    def h3(self, text):
        par = self.doc.add_paragraph(style="Heading 3")
        self._runs(par, text)
        return par

    # ---------------------------------------------------------------- списки
    def _list_par(self, text, num_id, keep_next=False, align=A.JUSTIFY):
        par = self.doc.add_paragraph(style="List Paragraph")
        ppr = par._p.get_or_add_pPr()
        numPr = ppr.get_or_add_numPr()
        numPr.get_or_add_ilvl().val = 0
        numPr.get_or_add_numId().val = num_id
        pf = par.paragraph_format
        pf.alignment = align
        pf.left_indent = Pt(self.P.list_left / 20)
        pf.first_line_indent = Pt(-self.P.list_hanging / 20)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.5
        pf.keep_with_next = keep_next
        self._runs(par, text)
        return par

    @staticmethod
    def _ends(items, end, last_end):
        out = []
        for i, it in enumerate(items):
            t = it.rstrip().rstrip(";.")
            out.append(t + (last_end if i == len(items) - 1 else end))
        return out

    def bullets(self, items, end=";", last_end="."):
        for t in self._ends(items, end, last_end):
            self._list_par(t, self.num_bullet)

    def numbered(self, items, end=";", last_end=".", align=A.JUSTIFY):
        nid = self._add_num(91, restart=True)
        for t in self._ends(items, end, last_end):
            self._list_par(t, nid, align=align)

    # ---------------------------------------------------------------- таблицы
    def _cell(self, cell, text, size, bold=None, align=A.LEFT, fill=None):
        cell.text = ""
        lines = str(text).split("\n")
        par = cell.paragraphs[0]
        for li, line in enumerate(lines):
            if li > 0:
                par = cell.add_paragraph()
            self._pf(par, align, 0, 0, 0, 1.0)
            self._runs(par, line, size, bold)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if fill:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), fill)
            tcPr.append(shd)

    def table(self, number, title, header, rows, widths_mm, size=None, align_cols=None, merge_rows=None,
              caption=True, header_fill=None, bold_rows=None, keep_together=None, head_rows=None,
              cant_split=True):
        P = self.P
        size = size or P.table
        if caption:
            cap = self.doc.add_paragraph()
            self._pf(cap, A.LEFT, 0, P.tcap_before, P.tcap_after, 1.5, True)
            self._runs(cap, f"Таблица {self.R(str(number))} {DASH} {title}")
        ncols = len(header)
        t = self.doc.add_table(rows=1, cols=ncols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        total_tw = int(sum(Mm(w).twips for w in widths_mm))
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:w"), str(total_tw))
        tblW.set(qn("w:type"), "dxa")
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), "000000")
            borders.append(b)
        tblPr.append(borders)
        mar = OxmlElement("w:tblCellMar")
        for k, v in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
            e = OxmlElement(f"w:{k}")
            e.set(qn("w:w"), str(v))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tblPr.append(mar)
        align_cols = align_cols or [A.LEFT] * ncols
        hdr = t.rows[0]
        trPr = hdr._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))
        trPr.append(OxmlElement("w:cantSplit"))
        for i, h in enumerate(header):
            self._cell(hdr.cells[i], h, size, True, A.CENTER, header_fill or P.header_fill)
        merge_rows = set(merge_rows or [])
        bold_rows = set(bold_rows or [])
        head_rows = set(head_rows or [])
        nrows = len(rows)
        for ri, row in enumerate(rows):
            r = t.add_row()
            if cant_split:
                r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if ri in merge_rows:
                m = r.cells[0].merge(r.cells[ncols - 1])
                self._cell(m, row[0], size, True, A.CENTER, "F2F2F2")
                continue
            if ri in head_rows:
                for ci in range(ncols):
                    self._cell(r.cells[ci], row[ci] if ci < len(row) else "", size, True, A.CENTER,
                               header_fill or P.header_fill)
                continue
            for ci in range(ncols):
                val = row[ci] if ci < len(row) else ""
                self._cell(r.cells[ci], val, size, True if ri in bold_rows else None, align_cols[ci])
        grid = t._tbl.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths_mm):
            gc.set(qn("w:w"), str(int(Mm(w).twips)))
        for row in t.rows:
            seen = set()
            ci = 0
            for c in row.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                span = c._tc.grid_span
                c.width = Mm(sum(widths_mm[ci:ci + span]))
                ci += span
        kt = keep_together if keep_together is not None else nrows <= 5
        if kt:
            keep = t.rows[:-1]
        elif nrows >= 4:
            # длинная таблица переносится, но без «висячих» строк: шапка не остается с одной строкой
            # внизу страницы, последняя строка не уходит на новую страницу одна
            keep = [t.rows[0], t.rows[1], t.rows[-2]]
        else:
            keep = []
        for row in keep:
            for c in row.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.keep_with_next = True
        self.empty()
        return t

    # ---------------------------------------------------------------- рисунки
    def figure(self, number, title, path, width_mm=150):
        P = self.P
        par = self.doc.add_paragraph()
        self._pf(par, A.CENTER, 0, P.fig_before, 3, 1.0, True)
        par.add_run().add_picture(path, width=Mm(width_mm))
        cap = self.doc.add_paragraph()
        self._pf(cap, A.CENTER, 0, 0, 6, 1.5)
        self._runs(cap, f"Рисунок {self.R(str(number))} {DASH} {title}")
        return par

    def placeholder(self, number, title, lines, out_dir, key, width_mm=120, ratio=0.5):
        os.makedirs(out_dir, exist_ok=True)
        path = os.path.join(out_dir, f"{key}.png")
        make_placeholder(path, lines, ratio=ratio)
        return self.figure(number, title, path, width_mm)

    # ---------------------------------------------------------------- содержание и титул
    def toc_entry(self, text, page, level=1):
        par = self.doc.add_paragraph()
        left = self.P.toc_l2_indent_cm if level == 2 else 0
        self._pf(par, A.LEFT, 0, 0, 0, 1.5, left_cm=left)
        par.paragraph_format.tab_stops.add_tab_stop(Pt(TEXT_W_TW / 20), WD_TAB_ALIGNMENT.RIGHT,
                                                    WD_TAB_LEADER.DOTS)
        self._runs(par, f"{text}\t{page}")
        return par

    def title_line(self, text, bold=False, before=0, after=0, size=None, align=A.CENTER):
        par = self.doc.add_paragraph()
        self._pf(par, align, 0, before, after, 1.5)
        self._runs(par, text, size, bold)
        return par

    def save(self, path):
        normalize_order(self.doc)
        self.doc.save(path)
        return path


def make_placeholder(path, lines, width_px=1600, ratio=0.5, label="МЕСТО ДЛЯ ФОТО / СКРИНШОТА"):
    height_px = int(width_px * ratio)
    img = Image.new("RGB", (width_px, height_px), (242, 242, 242))
    d = ImageDraw.Draw(img)
    d.rectangle([4, 4, width_px - 5, height_px - 5], outline=(128, 128, 128), width=4)
    f1 = ImageFont.truetype(FONT_BOLD_PATH, 52)
    f2 = ImageFont.truetype(FONT_PATH, 40)
    y = height_px // 2 - 50 - 28 * len(lines)
    w = d.textlength(label, font=f1)
    d.text(((width_px - w) / 2, y), label, fill=(90, 90, 90), font=f1)
    y += 100
    for ln in lines:
        w = d.textlength(ln, font=f2)
        d.text(((width_px - w) / 2, y), ln, fill=(110, 110, 110), font=f2)
        y += 54
    img.save(path)
    return path


def make_resolver(order):
    nums = {}
    cf = ct = 0
    for kind, key in order:
        if kind == "fig":
            cf += 1
            nums[("fig", key)] = cf
        else:
            ct += 1
            nums[("tbl", key)] = ct

    def R(text):
        def rep(m):
            k = ("fig" if m.group(1) == "F" else "tbl", m.group(2))
            if k not in nums:
                raise KeyError(f"нет объекта {k}")
            return str(nums[k])
        return re.sub(r"\{([FT]):([a-z0-9_]+)\}", rep, text)

    return R, cf, ct
